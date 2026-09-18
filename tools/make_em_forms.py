#!/usr/bin/env python3
"""Render the Editorial Manager forms from their Markdown sources.

    python tools/make_em_forms.py

Writes into submission/:
  cover_letter.docx   from submission/cover_letter.md
  highlights.docx     from submission/highlights.md, with Elsevier's 85-character
                      per-bullet limit enforced

The Markdown files are the single source of truth; the Word files EM wants are
derived rather than maintained by hand, so the two can never disagree.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt
except ImportError:  # pragma: no cover - environment problem, not a logic branch
    sys.exit("[fail] needs python-docx: pip install python-docx")

ROOT = Path(__file__).resolve().parents[1]
SUB = ROOT / "submission"
HIGHLIGHT_LIMIT = 85


def _document() -> Document:
    doc = Document()
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(11)
    return doc


def _plain(block: str) -> str:
    block = re.sub(r"^#+\s*", "", block)
    block = re.sub(r"\*\*(.*?)\*\*", r"\1", block)
    return re.sub(r"\*(.*?)\*", r"\1", block)


def cover_letter() -> None:
    src, out = SUB / "cover_letter.md", SUB / "cover_letter.docx"
    doc = _document()
    for block in re.split(r"\n\s*\n", src.read_text(encoding="utf-8")):
        block = block.strip()
        if not block or set(block) <= set("=-"):
            continue
        clean = _plain(block)
        if clean.lstrip().startswith(("-", "1.", "2.", "3.")):
            for line in clean.splitlines():
                line = re.sub(r"^\s*(?:-|\d\.)\s*", "", line).strip()
                if line:
                    doc.add_paragraph(line, style="List Bullet")
        else:
            doc.add_paragraph(re.sub(r"\s*\n\s*", " ", clean))
    doc.save(out)
    print(f"[ok  ] {out.relative_to(ROOT)}")


def highlights() -> None:
    src, out = SUB / "highlights.md", SUB / "highlights.docx"
    bullets = [
        _plain(line[2:]).strip()
        for line in src.read_text(encoding="utf-8").splitlines()
        if line.startswith("- ")
    ]
    if not 3 <= len(bullets) <= 5:
        sys.exit(f"[fail] Elsevier wants 3-5 highlights, found {len(bullets)}")
    too_long = [b for b in bullets if len(b) > HIGHLIGHT_LIMIT]
    if too_long:
        sys.exit(f"[fail] over {HIGHLIGHT_LIMIT} characters: {too_long}")
    doc = _document()
    doc.add_paragraph("Highlights")
    for b in bullets:
        doc.add_paragraph(b, style="List Bullet")
    doc.save(out)
    print(f"[ok  ] {out.relative_to(ROOT)}: {len(bullets)} bullets, "
          f"longest {max(len(b) for b in bullets)} characters")


if __name__ == "__main__":
    cover_letter()
    highlights()
