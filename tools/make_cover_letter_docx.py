#!/usr/bin/env python3
"""Render submission/cover_letter.md as the .docx that Editorial Manager accepts.

The letter is written and reviewed as Markdown; Editorial Manager wants a Word
file, so the Word file is derived rather than maintained by hand.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt
except ImportError:  # pragma: no cover - environment problem, not a logic branch
    sys.exit("[fail] needs python-docx: pip install python-docx")

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "submission/cover_letter.md"
OUT = ROOT / "submission/cover_letter.docx"


def main() -> None:
    doc = Document()
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(11)
    for block in re.split(r"\n\s*\n", SRC.read_text(encoding="utf-8")):
        block = block.strip()
        if not block or set(block) <= set("=-"):
            continue
        block = re.sub(r"^#+\s*", "", block)
        clean = re.sub(r"\*\*(.*?)\*\*", r"\1", block)
        clean = re.sub(r"\*(.*?)\*", r"\1", clean)
        if clean.lstrip().startswith(("-", "1.", "2.", "3.")):
            for line in clean.splitlines():
                line = re.sub(r"^\s*(?:-|\d\.)\s*", "", line).strip()
                if line:
                    doc.add_paragraph(line, style="List Bullet")
        else:
            doc.add_paragraph(re.sub(r"\s*\n\s*", " ", clean))
    doc.save(OUT)
    print(f"[ok  ] {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
